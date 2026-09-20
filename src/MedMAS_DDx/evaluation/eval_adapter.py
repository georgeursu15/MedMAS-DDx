#!/usr/bin/env python


from __future__ import annotations

import argparse
import csv
from datetime import datetime
from pathlib import Path
from typing import Any

from MedMAS_DDx.evaluation.eval_runner import (
    DEFAULT_CSV,
    DEFAULT_OUTPUT_DIR,
    run_evaluation,
)

from MedMAS_DDx.evaluation import batch_runner as _batch

DEFAULT_BENCHMARK_CSV = Path(__file__).resolve().parents[3] / "Benchmark" / "MedQDx_Benchmark.csv"



_P1_BLOCK = [
    "Sys_P1_Top1",       
    "Sys_P1_Top1_OK",    
    "Sys_P1_Top3_OK",   
    "Sys_P1_SimTop1",    
    "Sys_P1_SimTop3",    
]

_P2_BLOCK = [
    "Sys_P2_Top1",       
    "Sys_P2_Top1_OK",
    "Sys_P2_Top3_OK",
    "Sys_P2_SimTop1",
    "Sys_P2_SimTop3",
    "Sys_FollowupQs",    
    "Sys_Status",        
]

_QUESTION_PREFIXES = ("question_", "answer_", "diagnosis_", "similarity_")


def _result_to_columns(res: dict[str, Any]) -> dict[str, Any]:
    return {
        "Sys_P1_Top1": res.get("phase1_top1") or "",
        "Sys_P1_Top1_OK": bool(res.get("top1_correct_phase1")),
        "Sys_P1_Top3_OK": bool(res.get("top3_correct_phase1")),
        "Sys_P1_SimTop1": round(res.get("sim_score_phase1_top1", 0.0), 4),
        "Sys_P1_SimTop3": round(res.get("sim_score_phase1_top3", 0.0), 4),
        "Sys_P2_Top1": res.get("phase2_top1") or "",
        "Sys_P2_Top1_OK": bool(res.get("top1_correct_phase2")),
        "Sys_P2_Top3_OK": bool(res.get("top3_correct_phase2")),
        "Sys_P2_SimTop1": round(res.get("sim_score_phase2_top1", 0.0), 4),
        "Sys_P2_SimTop3": round(res.get("sim_score_phase2_top3", 0.0), 4),
        "Sys_FollowupQs": res.get("followup_questions_count", 0),
        "Sys_Status": res.get("status", "ok"),
    }



def _read_csv_rows(csv_path: str | Path) -> tuple[list[str], list[dict[str, str]]]:
    csv_path = Path(csv_path)
    if not csv_path.exists():
        raise FileNotFoundError(f"CSV not found: {csv_path}")
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = list(reader.fieldnames or [])
        rows = [dict(r) for r in reader]
    return fieldnames, rows


def _plan_columns(
    original_fieldnames: list[str],
    before_block: list[str],
    after_block: list[str],
) -> tuple[list[str], list[str]]:
  
    lowered = [fn.lower() for fn in original_fieldnames]
    q_idxs = [
        i for i, fn in enumerate(lowered)
        if fn.startswith(_QUESTION_PREFIXES)
    ]

    if not q_idxs:
        header = list(original_fieldnames) + before_block + after_block
        return header, before_block + after_block

    q_start, q_end = q_idxs[0], q_idxs[-1]
    header = (
        original_fieldnames[:q_start]
        + before_block
        + original_fieldnames[q_start:q_end + 1]
        + after_block
        + original_fieldnames[q_end + 1:]
    )
    return header, before_block + after_block



def _autosize(ws, max_width: int = 60) -> None:
    from openpyxl.utils import get_column_letter

    for col_idx, col_cells in enumerate(ws.columns, start=1):
        longest = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            max(10, longest + 2), max_width
        )


def _write_xlsx(
    output_path: Path,
    header: list[str],
    original_fieldnames: list[str],
    system_columns: list[str],
    original_rows: list[dict[str, str]],
    results_by_caseid: dict[int, dict[str, Any]],
    *,
    sheet_title: str = "results",
) -> None:
   
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "openpyxl is required to write the Excel report.\n"
            "Install it with:  uv add openpyxl   (or: pip install openpyxl)"
        ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)

    sys_set = set(system_columns)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(header)

    orig_fill = PatternFill("solid", fgColor="DCE6F1")  
    sys_fill = PatternFill("solid", fgColor="E2EFDA")   
    bold = Font(bold=True)
    for col_idx, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = bold
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.fill = sys_fill if name in sys_set else orig_fill

    for row_idx, orig in enumerate(original_rows, start=1):
        res = results_by_caseid.get(row_idx, {})
        record = [
            (res.get(name, "") if name in sys_set else orig.get(name, ""))
            for name in header
        ]
        ws.append(record)

    ws.freeze_panes = "B2"
    _autosize(ws)
    wb.save(output_path)



def _register_test(
    batch_path: Path,
    batch_id: str,
    cfg: dict[str, Any],
    csv_path: str,
) -> str:
   
    base = _batch._make_test_id(cfg.get("prognosis"), cfg.get("variant"), 0)
    existing = _batch._read_batch_file(batch_path) if batch_path.exists() else None

    if not existing:
        _batch._create_batch_file(batch_path, batch_id, [cfg], [base], csv_path)
        return base

    used = {t.get("test_id") for t in existing.get("tests", [])}
    tid, n = base, 1
    while tid in used:
        n += 1
        tid = f"{base}_{n}"

    existing.setdefault("tests", []).append({
        "test_id": tid,
        "prognosis_filter": cfg.get("prognosis"),
        "variant_filter": cfg.get("variant"),
        "limit": cfg.get("limit"),
        "status": "pending",
        "error": None,
        "started_at": None,
        "completed_at": None,
        "total_cases": 0,
        "summary": {},
        "cases": [],
    })
    tests = existing["tests"]
    meta = existing.setdefault("meta", {})
    meta["total_tests"] = len(tests)
    meta["pending_tests"] = sum(
        1 for t in tests if t.get("status") in ("pending", "running")
    )
    meta["status"] = "in_progress"
    _batch._write_batch_file(batch_path, existing)
    return tid


def _classify_run(summary: dict[str, Any]) -> tuple[str, str | None]:
    """Mirror batch_runner's ok / partial / failed classification."""
    total_eval = summary.get("total_evaluated", 0)
    total_err = summary.get("total_errors", 0)
    if total_eval == 0:
        return "failed", "all cases failed"
    if total_err > 0:
        return "partial", f"{total_err} case error(s)"
    return "ok", None


def run_adapter(
    csv_path: str | Path,
    *,
    insert_around_questions: bool,
    variant: str | None = "50",
    prognosis_filter: str | None = None,
    limit: int | None = None,
    output_path: str | Path | None = None,
    batch_id: str | None = None,
    use_patient_llm: bool = False,
    patient_llm_model: str | None = None,
    log_dir: Path | None = None,
) -> Path:
  
    original_fieldnames, original_rows = _read_csv_rows(csv_path)

    if insert_around_questions:
        before_block, after_block = _P1_BLOCK, _P2_BLOCK
    else:
        before_block, after_block = [], _P1_BLOCK + _P2_BLOCK

    header, system_columns = _plan_columns(
        original_fieldnames, before_block, after_block
    )

    batch_id = batch_id or _batch._make_batch_id()
    batch_dir = DEFAULT_OUTPUT_DIR / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)
    batch_path = batch_dir / _batch.BATCH_FILENAME      
    run_path = batch_dir / "run.txt"

    cfg = {"prognosis": prognosis_filter, "variant": variant, "limit": limit}
    test_id = _register_test(batch_path, batch_id, cfg, str(csv_path))
    report_path = batch_dir / f"{test_id}.json"

    print(f"\n[Adapter] Batch ID: {batch_id}")
    print(f"[Adapter] Test ID : {test_id}")
    print(f"[Adapter] Output folder: {batch_dir}\n")

    run_fresh = not run_path.exists()
    with run_path.open("a", encoding="utf-8") as run_log:
        if run_fresh:
            run_log.write(f"▶ batch {batch_id}\n\n")
        run_log.write(f"── Starting: {test_id}\n")
        run_log.flush()

        _batch._update_test_entry(
            batch_path, test_id,
            status="running", started_at=datetime.now().isoformat(),
        )

        report = run_evaluation(
            csv_path=csv_path,
            variant_filter=variant,
            prognosis_filter=prognosis_filter,
            limit=limit,
            log_dir=log_dir,
            batch_id=batch_id,
            output_path=report_path,     
            use_patient_llm=use_patient_llm,
            patient_llm_model=patient_llm_model,
        )

        status, err = _classify_run(report.get("summary", {}))
        _batch._update_test_entry(
            batch_path, test_id,
            status=status, completed_at=datetime.now().isoformat(),
            test_result=report, error=err,
        )
        mark = "✓ done" if status == "ok" else f"✗ {status}"
        run_log.write(f"── {mark}: {test_id}" + (f"  ({err})" if err else "") + "\n\n")

    results_by_caseid: dict[int, dict[str, Any]] = {
        res["case_id"]: _result_to_columns(res)
        for res in report.get("cases", [])
    }

    if output_path is None:
        stem = Path(csv_path).stem
        output_path = batch_dir / f"{stem}_with_metrics_{test_id}.xlsx"
    output_path = Path(output_path)

    _write_xlsx(
        output_path,
        header,
        original_fieldnames,
        system_columns,
        original_rows,
        results_by_caseid,
        sheet_title="evaluation",
    )

    print(f"\n[Adapter] Batch folder : {batch_dir}")
    print(f"[Adapter]   batch.json  : {batch_path.name}")
    print(f"[Adapter]   run.txt     : {run_path.name}")
    print(f"[Adapter]   report json : {report_path.name}")
    print(f"[Adapter]   excel       : {output_path.name}")
    print(f"[Adapter] Rows: {len(original_rows)} | metrics filled: {len(results_by_caseid)}\n")
    return output_path



def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "MedMAS_DDx evaluation adapter — runs the existing two-phase "
            "evaluation and writes the metrics into an Excel workbook."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument(
        "--mode", choices=["medqdx", "normal"], required=True,
        help=(
            "medqdx = insert Phase 1 metrics before the question columns and "
            "Phase 2 metrics after them (default input: MedQDx_Benchmark.csv).  "
            "normal = append metrics at the end (default input: PatientCases.csv)."
        ),
    )
    parser.add_argument(
        "--csv", type=str, default=None,
        help="Input CSV.  Defaults to MedQDx_Benchmark.csv (medqdx) or PatientCases.csv (normal).",
    )
    parser.add_argument(
        "--variant", type=str, choices=["100", "80", "50"], default="50",
        help="Case variant to evaluate. Default: 50 (the MedQDx partial-case protocol).",
    )
    parser.add_argument("--prognosis", type=str, default=None,
                        help="Filter by expected prognosis (case-insensitive substring).")
    parser.add_argument("--limit", type=int, default=None,
                        help="Maximum number of cases to evaluate (quick smoke tests).")
    parser.add_argument("--output", type=str, default=None,
                        help="Output .xlsx path. Defaults to inside the batch folder if omitted.")
    parser.add_argument("--batch-id", type=str, default=None, metavar="BATCH_ID",
                        help="Batch folder name under eval_results/. Auto-generated if omitted.")
    parser.add_argument("--log-dir", type=str, default=None,
                        help="Directory for per-case CrewAI log files. Default: logs/")
    parser.add_argument("--patient-llm", action="store_true",
                        help="Use the simulated-patient LLM to answer follow-ups (as before).")
    parser.add_argument("--patient-llm-model", type=str, default=None, metavar="MODEL",
                        help="Optional model override for the patient LLM.")
    args = parser.parse_args()

    if args.csv:
        csv_path: str | Path = args.csv
    elif args.mode == "medqdx":
        csv_path = DEFAULT_BENCHMARK_CSV
    else:
        csv_path = DEFAULT_CSV

    log_dir = (
        Path(args.log_dir) if args.log_dir
        else Path(__file__).resolve().parents[3] / "logs"
    )

    run_adapter(
        csv_path=csv_path,
        insert_around_questions=(args.mode == "medqdx"),
        variant=args.variant,
        prognosis_filter=args.prognosis,
        limit=args.limit,
        output_path=args.output,
        batch_id=args.batch_id,
        use_patient_llm=args.patient_llm,
        patient_llm_model=args.patient_llm_model,
        log_dir=log_dir,
    )


if __name__ == "__main__":
    main()
